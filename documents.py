from __future__ import annotations

from html import escape
from datetime import datetime
from pathlib import Path
from textwrap import shorten

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from openai_service import build_generation_prompt
from questionnaire import QUESTIONS


def _safe_filename(value: str) -> str:
    allowed = "abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789-_"
    normalized = "".join(ch if ch in allowed else "_" for ch in value)
    return normalized.strip("_") or "wedding_ceremony"


def _filename_hint(answers: dict[str, str], chat_id: int) -> str:
    names = answers.get("couple_names", "")
    return _safe_filename(shorten(names, width=40, placeholder="")) if names else f"chat_{chat_id}"


def _timestamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def _set_default_font(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)


def create_questionnaire_docx(answers: dict[str, str], output_dir: Path, chat_id: int) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    document = Document()
    _set_default_font(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Анкета для свадебной церемонии")
    title_run.bold = True
    title_run.font.size = Pt(20)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')}")

    current_section = None
    for question in QUESTIONS:
        if question.section != current_section:
            document.add_heading(question.section, level=1)
            current_section = question.section
        paragraph = document.add_paragraph()
        paragraph.add_run(question.prompt).bold = True
        document.add_paragraph(answers.get(question.key, "").strip() or "Не заполнено")

    output_path = output_dir / f"{_filename_hint(answers, chat_id)}_questionnaire_{_timestamp()}.docx"
    document.save(output_path)
    return output_path


def create_ceremony_prompt_txt(answers: dict[str, str], output_dir: Path, chat_id: int) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{_filename_hint(answers, chat_id)}_ceremony_prompt_{_timestamp()}.txt"
    output_path.write_text(build_generation_prompt(answers), encoding="utf-8")
    return output_path


def create_ceremony_script_txt(script_text: str, answers: dict[str, str], output_dir: Path, chat_id: int) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{_filename_hint(answers, chat_id)}_ceremony_script_{_timestamp()}.txt"
    output_path.write_text(script_text.strip() + "\n", encoding="utf-8")
    return output_path


def create_ceremony_script_docx(script_text: str, answers: dict[str, str], output_dir: Path, chat_id: int) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    document = Document()
    _set_default_font(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Текст свадебной церемонии")
    title_run.bold = True
    title_run.font.size = Pt(20)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"Пара: {answers.get('couple_names', 'Не указано')}")

    for block in _script_blocks(script_text):
        if _looks_like_heading(block):
            document.add_heading(_clean_heading(block), level=1)
        else:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.add_run(block)

    output_path = output_dir / f"{_filename_hint(answers, chat_id)}_ceremony_text_{_timestamp()}.docx"
    document.save(output_path)
    return output_path


def create_ceremony_script_html(script_text: str, answers: dict[str, str], output_dir: Path, chat_id: int) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    couple_names = answers.get("couple_names", "Пара").strip() or "Пара"
    blocks = []
    for block in _script_blocks(script_text):
        if _looks_like_heading(block):
            blocks.append(f"<section class=\"block\"><h2>{escape(_clean_heading(block))}</h2></section>")
        else:
            blocks.append(f"<section class=\"block\"><p>{escape(block).replace(chr(10), '<br>')}</p></section>")

    html = f"""<!doctype html>
<html lang="ru">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Текст церемонии - {escape(couple_names)}</title>
  <style>
    body {{
      margin: 0;
      font-family: Arial, sans-serif;
      color: #1f2933;
      background: #f6f4ef;
      line-height: 1.6;
    }}
    main {{
      max-width: 900px;
      margin: 0 auto;
      padding: 40px 24px;
    }}
    header {{
      margin-bottom: 28px;
      border-bottom: 2px solid #c9b99a;
      padding-bottom: 18px;
    }}
    h1 {{
      margin: 0 0 8px;
      font-size: 32px;
      line-height: 1.2;
    }}
    .meta {{
      color: #52606d;
      font-size: 15px;
    }}
    .block {{
      background: #ffffff;
      border: 1px solid #ddd6c8;
      border-left: 5px solid #9f7f45;
      border-radius: 6px;
      margin: 16px 0;
      padding: 18px 20px;
      page-break-inside: avoid;
    }}
    h2 {{
      margin: 0;
      font-size: 22px;
      color: #7b5e2b;
    }}
    p {{
      margin: 0;
      white-space: normal;
    }}
    @media print {{
      body {{ background: #ffffff; }}
      main {{ padding: 0; }}
      .block {{ border-color: #cccccc; }}
    }}
  </style>
</head>
<body>
  <main>
    <header>
      <h1>Текст свадебной церемонии</h1>
      <div class="meta">Пара: {escape(couple_names)}<br>Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')}</div>
    </header>
    {''.join(blocks)}
  </main>
</body>
</html>
"""
    output_path = output_dir / f"{_filename_hint(answers, chat_id)}_ceremony_text_{_timestamp()}.html"
    output_path.write_text(html, encoding="utf-8")
    return output_path


def _script_blocks(script_text: str) -> list[str]:
    blocks = [block.strip() for block in script_text.strip().split("\n\n")]
    return [block for block in blocks if block]


def _looks_like_heading(block: str) -> bool:
    first_line = block.splitlines()[0].strip()
    return len(block.splitlines()) == 1 and (
        first_line.startswith("#")
        or first_line.endswith(":")
        or first_line.lower().startswith(("начало", "финал", "выход", "история", "кольца", "клятвы"))
    )


def _clean_heading(block: str) -> str:
    return block.strip().lstrip("#").strip().rstrip(":")
